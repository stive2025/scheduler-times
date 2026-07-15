"""
Exportación a Excel de attendance_daily enriquecida con el horario oficial
esperado (Horario A/B/C/D, Medio Tiempo, Especial) para cada registro.

Las horas oficiales deben mantenerse en sincronía con HORARIOS_DEFINIDOS
en Generar_informe_3.py.
"""

import io
from datetime import time

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .database import get_engine

HORARIOS_ESPERADOS = {
    "Horario A": {
        "entrada": time(8, 0), "salida_almuerzo": time(13, 0),
        "entrada_tarde": time(14, 0), "salida": time(17, 0),
    },
    "Horario B": {
        "entrada": time(9, 30), "salida_almuerzo": time(14, 0),
        "entrada_tarde": time(15, 0), "salida": time(18, 30),
    },
    "Horario C": {
        "entrada": time(7, 0), "salida_almuerzo": time(12, 0),
        "entrada_tarde": time(13, 30), "salida": time(16, 30),
    },
    "Horario D": {
        "entrada": time(9, 0), "salida_almuerzo": time(13, 30),
        "entrada_tarde": time(13, 30), "salida": time(15, 0),
    },
    "Horario Especial": {
        "entrada": time(8, 0), "salida_almuerzo": time(13, 0),
        "entrada_tarde": None, "salida": time(13, 0),
    },
}

MEDIO_TIEMPO_MANANA = {
    "entrada": time(8, 0), "salida_almuerzo": time(12, 0),
    "entrada_tarde": None, "salida": None,
}
MEDIO_TIEMPO_TARDE = {
    "entrada": None, "salida_almuerzo": None,
    "entrada_tarde": time(14, 0), "salida": time(17, 0),
}

_EXPORT_QUERY = """
    SELECT
        employee_id, employee_name, date,
        schedule_name, shift_type, shift_detail, shift_str,
        entry_morning, exit_morning, entry_afternoon, exit_afternoon,
        hours_morning, hours_afternoon, hours_total,
        delay_morning_min, delay_afternoon_min, delay_total_min,
        early_exit_min, overtime_hours, event, annotations, month_period
    FROM attendance_daily
    ORDER BY date DESC, employee_name ASC
"""

_COLUMN_ORDER = [
    "employee_id", "employee_name", "date",
    "schedule_name", "shift_type", "shift_detail", "shift_str",
    "entry_morning", "entrada_esperada",
    "exit_morning", "salida_almuerzo_esperada",
    "entry_afternoon", "entrada_tarde_esperada",
    "exit_afternoon", "salida_esperada",
    "hours_morning", "hours_afternoon", "hours_total",
    "delay_morning_min", "delay_afternoon_min", "delay_total_min",
    "early_exit_min", "overtime_hours", "event", "annotations", "month_period",
]

_HEADERS_ES = {
    "employee_id": "ID Empleado",
    "employee_name": "Nombre",
    "date": "Fecha",
    "schedule_name": "Horario Asignado",
    "shift_type": "Jornada",
    "shift_detail": "Turno",
    "shift_str": "Horario (Texto)",
    "entry_morning": "Entrada Mañana (Real)",
    "entrada_esperada": "Entrada Mañana (Esperada)",
    "exit_morning": "Salida Almuerzo (Real)",
    "salida_almuerzo_esperada": "Salida Almuerzo (Esperada)",
    "entry_afternoon": "Entrada Tarde (Real)",
    "entrada_tarde_esperada": "Entrada Tarde (Esperada)",
    "exit_afternoon": "Salida (Real)",
    "salida_esperada": "Salida (Esperada)",
    "hours_morning": "Horas Mañana",
    "hours_afternoon": "Horas Tarde",
    "hours_total": "Horas Totales",
    "delay_morning_min": "Atraso Mañana (min)",
    "delay_afternoon_min": "Atraso Tarde (min)",
    "delay_total_min": "Atraso Total (min)",
    "early_exit_min": "Salida Anticipada (min)",
    "overtime_hours": "Horas Extra",
    "event": "Evento",
    "annotations": "Anotaciones",
    "month_period": "Periodo",
}


def _horario_esperado(schedule_name: str, shift_detail: str) -> dict | None:
    name = (schedule_name or "").strip()
    if name == "Medio Tiempo":
        if (shift_detail or "").strip() == "tarde":
            return MEDIO_TIEMPO_TARDE
        return MEDIO_TIEMPO_MANANA
    return HORARIOS_ESPERADOS.get(name)


def _fmt_time(value) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value)[:5]


def _build_dataframe() -> pd.DataFrame:
    engine = get_engine()
    df = pd.read_sql(_EXPORT_QUERY, engine)

    if df.empty:
        return pd.DataFrame(columns=[_HEADERS_ES[c] for c in _COLUMN_ORDER])

    entrada_esp, salida_alm_esp, entrada_tar_esp, salida_esp = [], [], [], []
    for _, row in df.iterrows():
        h = _horario_esperado(row["schedule_name"], row["shift_detail"])
        if h:
            entrada_esp.append(_fmt_time(h["entrada"]))
            salida_alm_esp.append(_fmt_time(h["salida_almuerzo"]))
            entrada_tar_esp.append(_fmt_time(h["entrada_tarde"]))
            salida_esp.append(_fmt_time(h["salida"]))
        else:
            entrada_esp.append(None)
            salida_alm_esp.append(None)
            entrada_tar_esp.append(None)
            salida_esp.append(None)

    df["entrada_esperada"] = entrada_esp
    df["salida_almuerzo_esperada"] = salida_alm_esp
    df["entrada_tarde_esperada"] = entrada_tar_esp
    df["salida_esperada"] = salida_esp

    for col in ("entry_morning", "exit_morning", "entry_afternoon", "exit_afternoon"):
        df[col] = df[col].apply(_fmt_time)

    df["date"] = pd.to_datetime(df["date"], errors="coerce").dt.strftime("%Y-%m-%d")

    df = df[_COLUMN_ORDER]
    df = df.rename(columns=_HEADERS_ES)
    return df


def build_attendance_export() -> bytes:
    """Genera el archivo Excel completo de attendance_daily + horario esperado."""
    df = _build_dataframe()

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        sheet_name = "Asistencias"
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

            max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str).values]) if len(df) else len(str(col_name))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 32)

        ws.freeze_panes = "A2"

    buffer.seek(0)
    return buffer.getvalue()
